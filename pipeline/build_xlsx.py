import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

df = pd.read_pickle('/home/claude/dbtr/kap_db.pkl')
rest = pd.read_pickle('/home/claude/dbtr/restatements.pkl')
q3c = pd.read_pickle('/home/claude/dbtr/q3c.pkl')
OUT = '/home/claude/dbtr/DBTR_KAP_DB.xlsx'

F = 'Arial'
fn = Font(name=F, size=9); fb = Font(name=F, size=9, bold=True); fh = Font(name=F, size=9, bold=True, color='FFFFFF')
fgrey = Font(name=F, size=8, color='808080'); fgreen = Font(name=F, size=9, color='008000')
hfill = PatternFill('solid', fgColor='1F3864'); sfill = PatternFill('solid', fgColor='D9E1F2'); yfill = PatternFill('solid', fgColor='FFF2CC')
NUM = '#,##0.0;-#,##0.0;"-"'; NUM0 = '#,##0;-#,##0;"-"'; DT = 'dd.mm.yyyy'

wb = Workbook()

# ---------------- README ----------------
ws = wb.active; ws.title = 'README'
notes = [
 ('DEUTSCHE BANK A.Ş. – KAP Banka Finansal Rapor (Konsolide Olmayan) veri tabanı', fb),
 ('Kaynak: KAP bildirimleri 977767 … 1649170 (2021Ç3 – 2026Ç2, 20 bildirim). Her bildirimde Cari + Önceki dönem kolonları ayrıştırıldı; böylece seri 31.12.2020 / 30.09.2020\'ye kadar geriye gider.', fn),
 ('Birim: Tüm değerler "value_mn" kolonunda milyon TL\'ye normalize edilmiştir (2026Ç2 bildirimi 1.000.000 TL, öncekiler 1.000 TL birimiyle yayımlandı). Hisse Başına Kar TL/hisse olarak bildirildiği gibi bırakıldı, ölçeklenmedi.', fn),
 ('Enflasyon muhasebesi: Hiçbir dönemde "Net Parasal Pozisyon Karı" ≠ 0 değil → TMS 29 uygulanmamış, tüm seriler nominal TL.', fn),
 ('', fn),
 ('SAYFALAR', fb),
 ('DB         : Uzun format ham veri tabanı. Her satır = bildirim × tablo × kalem × kolon. role=Cari/Önceki, basis=POINT (bilanço) / YTD (kümülatif) / Q3M (bildirilen 3 aylık). use=1 → o dönem için tercih edilen gözlem (Cari varsa Cari, yoksa sonraki bildirimdeki Önceki).', fn),
 ('Items      : Seri sözlüğü (series_key = tablo|para birimi|hiyerarşik yol). Hiyerarşi PDF\'teki girinti seviyesinden (level) türetildi.', fn),
 ('BS_TOT/TP/YP, OBS_TOT/TP/YP : Bilanço ve nazım hesaplar, dönem sonu stok. Formül: SUMIFS(DB, series_key, p_end, basis, use=1). 31.03.2021 ve 30.06.2021 bilançoları bildirimlerde yok (Önceki dönem her zaman 31.12).', fn),
 ('PL_YTD, OCI_YTD, CF_YTD : Kümülatif (yılbaşından) akımlar.', fn),
 ('PL_Q, CF_Q : Türetilmiş çeyreklik akımlar = YTD − önceki YTD (Ç1 = YTD). 30.09.2020 kolonu bildirilen 3 aylık değerden.', fn),
 ('EQ         : Özkaynak değişim tablosu (satır = hareket × özkaynak bileşeni).', fn),
 ('Checks     : Formül bazlı tutarlılık kontrolleri (Aktif=Pasif, TP+YP=Toplam, K/Z net kar = Bilanço dönem karı, NII = faiz gelir + gider).', fn),
 ('Restatements : Sonraki bildirimdeki "Önceki Dönem" değeri ile orijinal "Cari Dönem" değerinin farkı. kategori=rounding → 2026Ç2 bildiriminin mn TL\'ye yuvarlamasından; restatement → gerçek yeniden sınıflama.', fn),
 ('Q3M_Check  : Bildirilen 3 aylık K/Z vs YTD farkından türetilen; HBK hariç tüm farklar yuvarlama kaynaklı.', fn),
 ('', fn),
 ('Yeni bildirim eklemek için: parse_kap.py → build_db.py → build_xlsx.py (DB sayfasına satır eklenince tüm geniş tablolar formülle güncellenir).', fgrey),
]
for i, (t, f) in enumerate(notes, 1):
    c = ws.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=False)
ws.column_dimensions['A'].width = 180

# ---------------- Filings ----------------
ws = wb.create_sheet('Filings')
fil = df.drop_duplicates('file_id')[['file_id', 'filing', 'year', 'quarter', 'pub_date', 'unit', 'audit_firm', 'audit_type', 'audit_result']].sort_values(['year', 'quarter'])
fil['kap_url'] = 'https://www.kap.org.tr/tr/Bildirim/' + fil.file_id
cnt = df.groupby('file_id').size().rename('db_rows')
fil = fil.merge(cnt, left_on='file_id', right_index=True)
hdr = ['file_id', 'filing', 'year', 'quarter', 'pub_date', 'unit (TL)', 'audit_firm', 'audit_type', 'audit_result', 'kap_url', 'db_rows']
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h); c.font = fh; c.fill = hfill
for i, r in enumerate(fil.itertuples(index=False), 2):
    for j, v in enumerate(r, 1):
        c = ws.cell(i, j, v if not (isinstance(v, float) and np.isnan(v)) else None); c.font = fn
for j, w in enumerate([10, 8, 6, 7, 11, 10, 60, 10, 10, 42, 8], 1):
    ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'

# ---------------- DB ----------------
ws = wb.create_sheet('DB')
cols = ['file_id', 'filing', 'stmt', 'ccy', 'basis', 'role', 'p_start', 'p_end', 'eq_component', 'level', 'label', 'path', 'series_key', 'note', 'raw', 'value', 'unit', 'value_mn', 'value_unit', 'use', 'col', 'page']
dbo = df.sort_values(['year', 'quarter', 'stmt', 'row_order', 'col']).reset_index(drop=True)
dbo = dbo[dbo.value.notna()].reset_index(drop=True)   # keep only cells with a number
for j, h in enumerate(cols, 1):
    c = ws.cell(1, j, h); c.font = fh; c.fill = hfill
rows = dbo[cols].to_dict('records')
for i, r in enumerate(rows, 2):
    for j, h in enumerate(cols, 1):
        v = r[h]
        if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
        if isinstance(v, float) and np.isnan(v): v = None
        if v is pd.NaT or v is None: continue
        c = ws.cell(i, j, v)
        if h in ('p_start', 'p_end'): c.number_format = DT
        if h == 'value_mn': c.number_format = '#,##0.000'
NDB = len(rows) + 1
tab = Table(displayName='tblDB', ref=f'A1:{L(len(cols))}{NDB}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
ws.add_table(tab)
for j, w in enumerate([8, 7, 5, 7, 6, 7, 10, 10, 18, 5, 40, 60, 60, 7, 12, 12, 8, 12, 8, 4, 30, 5], 1):
    ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'
for row in ws.iter_rows(min_row=2, max_row=NDB):
    for c in row: c.font = fn
ws.sheet_properties.tabColor = '808080'

# column letters in DB
CI = {h: L(j) for j, h in enumerate(cols, 1)}
def rng(h): return f"DB!${CI[h]}$2:${CI[h]}${NDB}"

# ---------------- Items ----------------
u = df[df.use == 1]
items = (u.sort_values(['year', 'quarter', 'stmt', 'row_order'])
          .drop_duplicates('series_key')[['stmt', 'ccy', 'eq_component', 'level', 'label', 'path', 'series_key']])
# preserve the ordering of the most recent filing for each statement, append older-only items after their parent order
def item_order(st):
    latest = df[(df.stmt == st) & (df.file_id == '1649170')].drop_duplicates('path')[['path', 'row_order']]
    order = dict(zip(latest.path, latest.row_order))
    sub = items[items.stmt == st].copy()
    sub['o'] = sub.path.map(order)
    # unknown paths: place after the parent path's order
    for idx, r in sub[sub.o.isna()].iterrows():
        parent = ' > '.join(r.path.split(' > ')[:-1])
        sub.loc[idx, 'o'] = order.get(parent, 0) + 0.5
    return sub.sort_values(['o', 'ccy', 'eq_component'])
items = pd.concat([item_order(s) for s in ['BS', 'OBS', 'PL', 'OCI', 'CF', 'EQ']])
ws = wb.create_sheet('Items')
hdr = ['stmt', 'ccy', 'eq_component', 'level', 'label', 'path', 'series_key', 'n_obs(use=1)']
nobs = u.groupby('series_key').size()
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h); c.font = fh; c.fill = hfill
for i, r in enumerate(items.itertuples(index=False), 2):
    vals = [r.stmt, r.ccy, r.eq_component, r.level, r.label, r.path, r.series_key, int(nobs.get(r.series_key, 0))]
    for j, v in enumerate(vals, 1):
        ws.cell(i, j, v).font = fn
for j, w in enumerate([5, 7, 20, 5, 50, 70, 70, 8], 1):
    ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'

# ---------------- wide sheets ----------------
def periods_for(st, basis):
    d = u[(u.stmt == st) & (u.basis == basis)]
    return sorted(d.p_end.dropna().unique())

def wide_sheet(name, st, ccy, basis, title, derive_q=False, src=None):
    ws = wb.create_sheet(name)
    per = periods_for(st, 'YTD' if derive_q else basis)
    if st in ('BS', 'OBS'):
        sub = items[items.stmt == st].drop_duplicates('path').copy()
        sub['series_key'] = st + '|' + ccy + '|' + sub.path
    else:
        sub = items[(items.stmt == st) & (items.ccy == ccy)]
    ws['A1'] = title; ws['A1'].font = Font(name=F, size=11, bold=True)
    ws['A2'] = 'mn TL (nominal). ' + ('Türetilmiş çeyreklik akım: Ç1 = YTD, diğer = YTD − önceki YTD; ilk kolon bildirilen 3 aylık.' if derive_q else
                                     ('Dönem sonu stok.' if basis == 'POINT' else 'Yılbaşından kümülatif.')); ws['A2'].font = fgrey
    hdrs = ['series_key', 'Kalem', 'lvl']
    for j, h in enumerate(hdrs, 1):
        c = ws.cell(4, j, h); c.font = fh; c.fill = hfill
    ws.cell(3, 4, 'Dönem sonu →').font = fgrey
    for j, p in enumerate(per, 4):
        c = ws.cell(4, j, pd.Timestamp(p).to_pydatetime()); c.number_format = 'mm/yy'; c.font = fh; c.fill = hfill; c.alignment = Alignment(horizontal='right')
        q = (pd.Timestamp(p).month - 1) // 3 + 1
        c2 = ws.cell(3, j, f'{pd.Timestamp(p).year}Ç{q}'); c2.font = fgrey; c2.alignment = Alignment(horizontal='right')
    r0 = 5
    for i, r in enumerate(sub.itertuples(index=False), r0):
        ws.cell(i, 1, r.series_key).font = fgrey
        c = ws.cell(i, 2, r.label); c.font = fb if r.level <= 1 else fn
        c.alignment = Alignment(indent=int(r.level))
        ws.cell(i, 3, int(r.level)).font = fgrey
        for j, p in enumerate(per, 4):
            col = L(j)
            if derive_q:
                if j == 4:
                    f = f'=SUMIFS({rng("value_mn")},{rng("series_key")},$A{i},{rng("p_end")},{src}!{col}$4,{rng("basis")},"Q3M",{rng("use")},1)'
                else:
                    pc = L(j - 1)
                    f = f'=IF(MONTH({src}!{col}$4)=3,{src}!{col}{i},{src}!{col}{i}-{src}!{pc}{i})'
            else:
                f = f'=SUMIFS({rng("value_mn")},{rng("series_key")},$A{i},{rng("p_end")},{col}$4,{rng("basis")},"{basis}",{rng("use")},1)'
            c = ws.cell(i, j, f); c.font = fb if r.level <= 1 else fn; c.number_format = NUM
        if r.label == 'Hisse Başına Kar':
            for j in range(4, 4 + len(per)): ws.cell(i, j).number_format = '0.0000'
    ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 62; ws.column_dimensions['C'].width = 3
    for j in range(4, 4 + len(per)): ws.column_dimensions[L(j)].width = 9.5
    ws.freeze_panes = 'D5'
    ws.sheet_view.zoomScale = 90
    return ws, per, r0, len(sub)

info = {}
info['BS_TOT'] = wide_sheet('BS_TOT', 'BS', 'Toplam', 'POINT', 'Bilanço – Toplam')
info['BS_TP'] = wide_sheet('BS_TP', 'BS', 'TP', 'POINT', 'Bilanço – TP')
info['BS_YP'] = wide_sheet('BS_YP', 'BS', 'YP', 'POINT', 'Bilanço – YP')
info['OBS_TOT'] = wide_sheet('OBS_TOT', 'OBS', 'Toplam', 'POINT', 'Nazım Hesaplar – Toplam')
info['OBS_TP'] = wide_sheet('OBS_TP', 'OBS', 'TP', 'POINT', 'Nazım Hesaplar – TP')
info['OBS_YP'] = wide_sheet('OBS_YP', 'OBS', 'YP', 'POINT', 'Nazım Hesaplar – YP')
info['PL_YTD'] = wide_sheet('PL_YTD', 'PL', 'Toplam', 'YTD', 'Kar veya Zarar Tablosu – YTD')
info['PL_Q'] = wide_sheet('PL_Q', 'PL', 'Toplam', 'YTD', 'Kar veya Zarar Tablosu – Çeyreklik (türetilmiş)', derive_q=True, src='PL_YTD')
info['OCI_YTD'] = wide_sheet('OCI_YTD', 'OCI', 'Toplam', 'YTD', 'Kar veya Zarar ve Diğer Kapsamlı Gelir – YTD')
info['CF_YTD'] = wide_sheet('CF_YTD', 'CF', 'Toplam', 'YTD', 'Nakit Akış Tablosu – YTD')
info['CF_Q'] = wide_sheet('CF_Q', 'CF', 'Toplam', 'YTD', 'Nakit Akış Tablosu – Çeyreklik (türetilmiş)', derive_q=True, src='CF_YTD')
# CF_Q first column: there is no Q3M for CF -> blank it
wsq = wb['CF_Q']; _, per, r0, n = info['CF_Q']
for i in range(r0, r0 + n): wsq.cell(i, 4).value = None
wsq['A2'] = 'mn TL (nominal). Türetilmiş çeyreklik akım: Ç1 = YTD, diğer = YTD − önceki YTD. 30.09.2020 için 3 aylık bildirilmediğinden boş.'

# ---------------- EQ ----------------
ws = wb.create_sheet('EQ')
per = periods_for('EQ', 'YTD')
sub = items[items.stmt == 'EQ']
ws['A1'] = 'Özkaynaklar Değişim Tablosu – YTD hareketler (satır = hareket × bileşen)'; ws['A1'].font = Font(name=F, size=11, bold=True)
ws['A2'] = 'mn TL. Her blok yılbaşından dönem sonuna kümülatif hareket; Dönem Başı/Dönem Sonu Bakiyeler stok.'; ws['A2'].font = fgrey
for j, h in enumerate(['series_key', 'Hareket', 'Bileşen'], 1):
    c = ws.cell(4, j, h); c.font = fh; c.fill = hfill
for j, p in enumerate(per, 4):
    c = ws.cell(4, j, pd.Timestamp(p).to_pydatetime()); c.number_format = 'mm/yy'; c.font = fh; c.fill = hfill; c.alignment = Alignment(horizontal='right')
for i, r in enumerate(sub.itertuples(index=False), 5):
    ws.cell(i, 1, r.series_key).font = fgrey
    ws.cell(i, 2, r.label).font = fn
    ws.cell(i, 3, r.eq_component).font = fn
    for j, p in enumerate(per, 4):
        col = L(j)
        f = f'=SUMIFS({rng("value_mn")},{rng("series_key")},$A{i},{rng("p_end")},{col}$4,{rng("basis")},"YTD",{rng("use")},1)'
        c = ws.cell(i, j, f); c.font = fn; c.number_format = NUM
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 48; ws.column_dimensions['C'].width = 40
for j in range(4, 4 + len(per)): ws.column_dimensions[L(j)].width = 9.5
ws.freeze_panes = 'D5'

# ---------------- Checks (formula) ----------------
ws = wb.create_sheet('Checks')
ws['A1'] = 'Tutarlılık kontrolleri (formül; sıfırdan farklı hücre = sorun)'; ws['A1'].font = Font(name=F, size=11, bold=True)
bs_ws, bs_per, bs_r0, bs_n = info['BS_TOT']
def find_row(sheetinfo, label_path_end):
    ws_, per_, r0_, n_ = sheetinfo
    for i in range(r0_, r0_ + n_):
        if ws_.cell(i, 1).value.endswith(label_path_end): return i
    raise KeyError(label_path_end)
rA = find_row(info['BS_TOT'], 'VARLIKLAR > VARLIKLAR TOPLAMI'); rLi = find_row(info['BS_TOT'], 'YÜKÜMLÜLÜKLER > YÜKÜMLÜLÜKLER TOPLAMI')
rNI_bs = find_row(info['BS_TOT'], 'Kar veya Zarar > Dönem Net Kâr veya Zararı')
pl_ws, pl_per, pl_r0, pl_n = info['PL_YTD']
rNI = find_row(info['PL_YTD'], '> DÖNEM NET KARI VEYA ZARARI'); rII = find_row(info['PL_YTD'], '> FAİZ GELİRLERİ'); rIE = find_row(info['PL_YTD'], '> FAİZ GİDERLERİ (-)'); rNII = find_row(info['PL_YTD'], '> NET FAİZ GELİRİ VEYA GİDERİ')
rEq = find_row(info['BS_TOT'], 'YÜKÜMLÜLÜKLER > ÖZKAYNAKLAR')
eq_ws = wb['EQ']; eq_per = per
# header
ws.cell(3, 1, 'Kontrol').font = fb
for j, p in enumerate(bs_per, 2):
    c = ws.cell(3, j, pd.Timestamp(p).to_pydatetime()); c.number_format = 'mm/yy'; c.font = fh; c.fill = hfill
checks_rows = [
    ('Bilanço: Aktif Toplamı − Pasif Toplamı', lambda j: f"=BS_TOT!{L(j+4)}{rA}-BS_TOT!{L(j+4)}{rLi}"),
    ('Bilanço: TP+YP−Toplam, |fark|>0.5 olan satır sayısı', lambda j: f"=SUMPRODUCT(--(ABS(BS_TP!{L(j+4)}{bs_r0}:{L(j+4)}{bs_r0+bs_n-1}+BS_YP!{L(j+4)}{bs_r0}:{L(j+4)}{bs_r0+bs_n-1}-BS_TOT!{L(j+4)}{bs_r0}:{L(j+4)}{bs_r0+bs_n-1})>0.5))"),
    ('K/Z YTD Dönem Net Karı − Bilanço Dönem Net Karı', lambda j: f"=IFERROR(INDEX(PL_YTD!$D${rNI}:$ZZ${rNI},MATCH(BS_TOT!{L(j+4)}$4,PL_YTD!$D$4:$ZZ$4,0))-BS_TOT!{L(j+4)}{rNI_bs},\"n/a\")"),
    ('K/Z YTD: Faiz Gelirleri + Faiz Giderleri − NII', lambda j: f"=IFERROR(INDEX(PL_YTD!$D${rII}:$ZZ${rII},MATCH(BS_TOT!{L(j+4)}$4,PL_YTD!$D$4:$ZZ$4,0))+INDEX(PL_YTD!$D${rIE}:$ZZ${rIE},MATCH(BS_TOT!{L(j+4)}$4,PL_YTD!$D$4:$ZZ$4,0))-INDEX(PL_YTD!$D${rNII}:$ZZ${rNII},MATCH(BS_TOT!{L(j+4)}$4,PL_YTD!$D$4:$ZZ$4,0)),\"n/a\")"),
]
# equity: EQ Dönem Sonu Bakiyeler Toplam Özkaynak vs BS Özkaynaklar
rEqEnd = None
for i in range(5, 5 + len(sub)):
    if eq_ws.cell(i, 2).value == 'Dönem Sonu Bakiyeler' and eq_ws.cell(i, 3).value == 'Toplam Özkaynak': rEqEnd = i
checks_rows.append(('Özkaynak Değişim Dönem Sonu Toplam − Bilanço Özkaynaklar', lambda j: f"=IFERROR(INDEX(EQ!$D${rEqEnd}:$ZZ${rEqEnd},MATCH(BS_TOT!{L(j+4)}$4,EQ!$D$4:$ZZ$4,0))-BS_TOT!{L(j+4)}{rEq},\"n/a\")"))
for i, (lab, fgen) in enumerate(checks_rows, 4):
    ws.cell(i, 1, lab).font = fn
    for j in range(len(bs_per)):
        c = ws.cell(i, j + 2, fgen(j)); c.font = fn; c.number_format = '#,##0.000;-#,##0.000;"-"'
ws.column_dimensions['A'].width = 60
for j in range(2, 2 + len(bs_per)): ws.column_dimensions[L(j)].width = 9
ws.freeze_panes = 'B4'

# ---------------- Restatements ----------------
ws = wb.create_sheet('Restatements')
rest = rest.copy()
later_unit = df.drop_duplicates('filing').set_index('filing').unit
rest['later_unit'] = rest.later_filing.map(later_unit)
rest['rel'] = rest['diff'].abs() / rest.orig_cari.abs().clip(lower=1)
rest['category'] = np.where((rest.later_unit == 1_000_000) & (rest['diff'].abs() < 2), 'rounding (mn TL bildirimi)',
                   np.where(rest.series_key.str.contains('Hisse Başına'), 'EPS (ölçek/baz farkı)', 'restatement / reclass'))
rest = rest.sort_values(['category', 'p_end', 'series_key'])
hdr = ['series_key', 'basis', 'p_end', 'orig_filing', 'orig_cari (mn)', 'later_filing', 'onceki_val (mn)', 'diff (mn)', 'rel', 'category']
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h); c.font = fh; c.fill = hfill
for i, r in enumerate(rest.itertuples(index=False), 2):
    vals = [r.series_key, r.basis, r.p_end.to_pydatetime(), r.orig_filing, r.orig_cari, r.later_filing, r.onceki_val, r.diff, r.rel, r.category]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v); c.font = fn
        if j == 3: c.number_format = DT
        if j in (5, 7, 8): c.number_format = '#,##0.000'
        if j == 9: c.number_format = '0.00%'
for j, w in enumerate([90, 6, 10, 9, 12, 9, 12, 10, 7, 26], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'

# ---------------- Q3M check ----------------
ws = wb.create_sheet('Q3M_Check')
q3c = q3c.sort_values(['p_end', 'series_key'])
hdr = ['series_key', 'p_end', 'reported_3m (mn)', 'derived_3m (mn)', 'diff (mn)']
for j, h in enumerate(hdr, 1):
    c = ws.cell(1, j, h); c.font = fh; c.fill = hfill
for i, r in enumerate(q3c.itertuples(index=False), 2):
    vals = [r.series_key, r.p_end.to_pydatetime(), r.reported_3m, r.derived_3m, r.diff]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v); c.font = fn
        if j == 2: c.number_format = DT
        if j >= 3: c.number_format = '#,##0.000'
for j, w in enumerate([90, 10, 14, 14, 10], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT, 'DB rows', NDB - 1, 'items', len(items))
