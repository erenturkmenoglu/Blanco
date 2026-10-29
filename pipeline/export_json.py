"""Export the recalculated workbook + footnote data + FX rates into a single JSON for the web app."""
import json, math, pandas as pd, numpy as np
from openpyxl import load_workbook
from datetime import datetime, date

XLSX = '/home/claude/dbtr/DBTR_KAP_DB.xlsx'
OUT = '/home/claude/blanco/frontend/public/data/dbtr.json'

# ---- FX: TCMB döviz alış, değerleme günü (bilanço tarihinde geçerli bülten) ----
FX = {
 '2020-09-30': (7.8080, 9.1281), '2020-12-31': (7.3405, 9.0079), '2021-03-31': (8.3260, 9.7741), '2021-06-30': (8.7052, 10.3645),
 '2021-09-30': (8.88, 10.29), '2021-12-31': (12.9775, 14.6823), '2022-03-31': (14.6458, 16.3086), '2022-06-30': (16.6690, 17.5221),
 '2022-09-30': (18.5038, 17.9232), '2022-12-31': (18.6983, 19.9349), '2023-03-31': (19.1460, 20.8021), '2023-06-30': (25.8231, 28.1540),
 '2023-09-30': (27.3767, 29.0305), '2023-12-31': (29.4382, 32.5739), '2024-03-31': (32.2854, 34.8023), '2024-06-30': (32.8262, 35.1284),
 '2024-09-30': (34.0900, 38.0180), '2024-12-31': (35.2233, 36.7429), '2025-03-31': (37.7656, 40.7019), '2025-06-30': (39.7424, 46.5526),
 '2025-09-30': (41.4984, 48.6479), '2025-12-31': (42.8623, 50.4532), '2026-03-31': (44.3841, 51.0236), '2026-06-30': (46.5551, 53.0950),
}
FX_SOURCES = {
 'default': 'TCMB döviz alış – geçici vergi/dönem sonu değerleme kuru (bbdas.com.tr sirkülerleri; ilgili tarihte geçerli bülten)',
 '2021-09-30': 'bigpara.hurriyet.com.tr TCMB kurları 30.09.2021 (2 ondalık; kapanış bülteni)',
 '2024-03-31': 'İSMMMO 15.04.2024 yabancı para değerleri (TCMB 31.03.2024)',
 '2025-12-31': 'esasdenetim.com 31.12.2025 değerleme kurları (TCMB)',
}

wb = load_workbook(XLSX, data_only=True)

def hdr_periods(ws, row=4, c0=1):
    out = {}
    for j in range(c0, ws.max_column + 1):
        v = ws.cell(row, j).value
        if isinstance(v, datetime): out[j] = v.strftime('%Y-%m-%d')
    return out

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)): return None
        return round(float(v), 6)
    return None

# ---- Kalemler ----
ws = wb['Kalemler']; pc = hdr_periods(ws)
periods = sorted(pc.values())
items = []
for i in range(5, ws.max_row + 1):
    code = ws.cell(i, 1).value
    if not code: continue
    items.append(dict(code=code, name=ws.cell(i, 2).value, section=ws.cell(i, 3).value, kind=ws.cell(i, 4).value, status=ws.cell(i, 5).value,
                      source=ws.cell(i, 6).value, note=ws.cell(i, 7).value or '', values={p: num(ws.cell(i, j).value) for j, p in pc.items()}))
# ---- Kalemler_Q ----
ws = wb['Kalemler_Q']; qc = hdr_periods(ws)
flows_q = {}
for i in range(5, ws.max_row + 1):
    code = ws.cell(i, 1).value; blk = ws.cell(i, 3).value
    if not code: continue
    d = flows_q.setdefault(code, dict(code=code, name=ws.cell(i, 2).value, q={}, ttm={}))
    d['q' if blk == 'Q' else 'ttm'] = {p: num(ws.cell(i, j).value) for j, p in qc.items()}
# ---- Oranlar ----
ws = wb['Oranlar']; oc = hdr_periods(ws)
ratios = []
for i in range(5, ws.max_row + 1):
    code = ws.cell(i, 1).value
    if not code: continue
    fmt = ws.cell(i, oc and min(oc)).number_format
    unit = 'pct' if '%' in fmt else ('x' if 'x' in fmt else 'mn')
    ratios.append(dict(code=code, name=ws.cell(i, 2).value, method=ws.cell(i, 3).value, formula=ws.cell(i, 4).value, unit=unit,
                       values={p: num(ws.cell(i, j).value) for j, p in oc.items()}))
# ---- statements (wide) ----
statements = {}
for sh in ['BS_TOT', 'BS_TP', 'BS_YP', 'OBS_TOT', 'OBS_TP', 'OBS_YP', 'PL_YTD', 'PL_Q', 'OCI_YTD', 'CF_YTD', 'CF_Q']:
    ws = wb[sh]; cc = hdr_periods(ws)
    rows = []
    for i in range(5, ws.max_row + 1):
        k = ws.cell(i, 1).value
        if not k: continue
        rows.append(dict(key=k, label=ws.cell(i, 2).value, level=ws.cell(i, 3).value, values={p: num(ws.cell(i, j).value) for j, p in cc.items()}))
    statements[sh] = dict(title=ws['A1'].value, note=ws['A2'].value, periods=sorted(cc.values()), rows=rows)
ws = wb['EQ']; cc = hdr_periods(ws)
rows = []
for i in range(5, ws.max_row + 1):
    k = ws.cell(i, 1).value
    if not k: continue
    rows.append(dict(key=k, label=ws.cell(i, 2).value, component=ws.cell(i, 3).value, values={p: num(ws.cell(i, j).value) for j, p in cc.items()}))
statements['EQ'] = dict(title=ws['A1'].value, note=ws['A2'].value, periods=sorted(cc.values()), rows=rows)
# ---- checks / restatements / filings ----
ws = wb['Checks']; cc = hdr_periods(ws, row=3)
checks = [dict(name=ws.cell(i, 1).value, values={p: num(ws.cell(i, j).value) for j, p in cc.items()}) for i in range(4, ws.max_row + 1) if ws.cell(i, 1).value]
ws = wb['Restatements']
rest = [dict(series_key=r[0], basis=r[1], p_end=r[2].strftime('%Y-%m-%d'), orig_filing=r[3], orig=num(r[4]), later_filing=r[5], later=num(r[6]), diff=num(r[7]), category=r[9])
        for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
ws = wb['Filings']
filings = [dict(file_id=str(r[0]), filing=r[1], year=r[2], quarter=r[3], pub_date=r[4], unit=r[5], audit_firm=r[6], audit_type=r[7], audit_result=r[8], url=r[9], rows=r[10])
           for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

# ---- dipnot ----
dn = pd.read_pickle('/home/claude/dbtr/dipnot.pkl')
dn['p'] = dn.p_end.dt.strftime('%Y-%m-%d')
NAMES = {
 'capital_adequacy_ratio': ('Sermaye yeterliliği oranı (SYR)', '%'), 'tier_1_ratio': ('Ana sermaye yeterliliği oranı', '%'), 'cet1_ratio': ('Çekirdek sermaye yeterliliği oranı', '%'),
 'risk_weighted_assets': ('Risk ağırlıklı varlıklar', 'mn TL'), 'regulatory_equity': ('Özkaynak (SYR hesabı)', 'mn TL'), 'cet1_capital': ('Çekirdek sermaye', 'mn TL'),
 'liquidity_coverage_ratio': ('Likidite karşılama oranı (TP+YP)', '%'), 'liquidity_coverage_ratio_fx': ('Likidite karşılama oranı (YP)', '%'),
 'hqla_stock': ('Yüksek kaliteli likit varlık stoku', 'mn TL'), 'net_cash_outflows': ('Net nakit çıkışları (30 gün)', 'mn TL'), 'nsfr': ('Net istikrarlı fonlama oranı', '%'),
 'loan_interest_tl': ('TL kredi ortalama faiz', '%'), 'loan_interest_usd': ('USD kredi ortalama faiz', '%'), 'loan_interest_eur': ('EUR kredi ortalama faiz', '%'),
 'deposit_interest_tl': ('TL mevduat (diğer) ortalama faiz', '%'), 'deposit_interest_usd': ('USD mevduat ortalama faiz', '%'), 'deposit_interest_eur': ('EUR mevduat ortalama faiz', '%'),
 'bank_deposit_interest_tl': ('TL bankalar mevduatı ortalama faiz', '%'), 'bank_deposit_interest_eur': ('EUR bankalar mevduatı ortalama faiz', '%'),
 'funds_borrowed_interest_tl': ('TL alınan krediler ortalama faiz', '%'), 'funds_borrowed_interest_eur': ('EUR alınan krediler ortalama faiz', '%'),
 'due_from_banks_interest_tl': ('TL bankalar (plasman) ortalama faiz', '%'), 'due_from_banks_interest_usd': ('USD bankalar (plasman) ortalama faiz', '%'),
 'securities_fvtpl_interest_tl': ('TL GUD K/Z menkul ortalama faiz', '%'), 'securities_fvoci_interest_tl': ('TL GUD DKG menkul ortalama faiz', '%'),
 'cash_cb_interest_tl': ('TL nakit/TCMB ortalama faiz', '%'), 'cash_cb_interest_eur': ('EUR nakit/TCMB ortalama faiz', '%'),
 'demand_deposits': ('Vadesiz mevduat (+7 gün ihbarlı)', 'mn TL'), 'time_deposits': ('Vadeli mevduat', 'mn TL'),
 'npl_amount': ('Takipteki krediler (donuk alacak)', 'mn TL'), 'loan_loss_provisions': ('Donuk alacak karşılığı', 'mn TL'),
 'branch_count': ('Şube sayısı', 'adet'), 'employee_count': ('Çalışan sayısı', 'adet'),
}
scal = dn[dn.bucket.isna()]
dip_scalars = []
for code, g in scal.groupby('code', sort=False):
    nm, un = NAMES.get(code, (code, g.unit.iloc[0]))
    dip_scalars.append(dict(code=code, name=nm, unit=un, source=g.table.iloc[0], files=sorted(set(g.source)), note=g.note.iloc[0] or '',
                            values={r.p: num(r.value) for r in g.itertuples()}))
buck = dn[dn.bucket.notna()]
tables = {}
for (code, p), g in buck.groupby(['code', 'p']):
    tables.setdefault(code, {})[p] = {r.bucket: num(r.value) for r in g.itertuples()}
TABLE_META = {
 'rm_': ('Kalan vade (likidite riski)', ['demand', '0_1m', '1_3m', '3_12m', '1_5y', 'over_5y', 'unallocated', 'total'], ['Vadesiz', '1 aya kadar', '1-3 ay', '3-12 ay', '1-5 yıl', '5 yıl+', 'Dağıtılamayan', 'Toplam']),
 'rp_': ('Yeniden fiyatlama (faize duyarlılık)', ['0_1m', '1_3m', '3_12m', '1_5y', 'over_5y', 'non_interest', 'total'], ['1 aya kadar', '1-3 ay', '3-12 ay', '1-5 yıl', '5 yıl+', 'Faizsiz', 'Toplam']),
 'dep_': ('Mevduat vade yapısı', ['demand', 'notice_7d', '0_1m', '1_3m', '3_6m', '6_12m', 'over_1y', 'accumulating', 'total'], ['Vadesiz', '7 gün ihbarlı', '1 aya kadar', '1-3 ay', '3-6 ay', '6 ay-1 yıl', '1 yıl+', 'Birikimli', 'Toplam']),
}
ROWN = {'rm_total_assets': 'Toplam varlıklar', 'rm_total_liabilities': 'Toplam yükümlülükler', 'rm_liquidity_gap': 'Likidite açığı', 'rm_loans': 'Verilen krediler', 'rm_other_deposits': 'Diğer mevduat',
        'rm_bank_deposits': 'Bankalar mevduatı', 'rm_funds_borrowed': 'Diğer mali kuruluşlardan sağlanan fonlar', 'rm_cash_cb': 'Nakit değerler ve TCMB', 'rm_due_from_banks': 'Bankalar',
        'rm_net_off_bs': 'Net bilanço dışı pozisyon', 'rm_deriv_receivables': 'Türev alacaklar', 'rm_deriv_payables': 'Türev borçlar',
        'rp_total_assets': 'Toplam varlıklar', 'rp_total_liabilities': 'Toplam yükümlülükler', 'rp_loans': 'Verilen krediler', 'rp_other_deposits': 'Diğer mevduat', 'rp_bank_deposits': 'Bankalar mevduatı',
        'rp_funds_borrowed': 'Diğer mali kuruluşlardan sağlanan fonlar', 'rp_on_bs_long': 'Bilanço içi uzun pozisyon', 'rp_on_bs_short': 'Bilanço içi kısa pozisyon', 'rp_total_position': 'Toplam pozisyon',
        'rp_off_bs_long': 'Nazım uzun pozisyon', 'rp_off_bs_short': 'Nazım kısa pozisyon',
        'dep_maturity_total': 'Toplam mevduat', 'dep_maturity_fx': 'Döviz tevdiat hesabı', 'dep_maturity_commercial': 'Ticari kuruluşlar mevduatı', 'dep_maturity_banks': 'Bankalar mevduatı'}
dip_tables = []
for pref, (title, bk, bl) in TABLE_META.items():
    rows = [dict(code=c, name=ROWN.get(c, c), periods=v) for c, v in tables.items() if c.startswith(pref)]
    dip_tables.append(dict(group=pref.strip('_'), title=title, buckets=bk, bucket_labels=bl, rows=rows))

out = dict(
    meta=dict(bank='Deutsche Bank A.Ş.', scope='Konsolide olmayan (solo)', generated=datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC'),
              unit='mn TL (nominal; TMS 29 uygulanmamış)', filings=filings,
              sources=['KAP Banka Finansal Rapor bildirimleri (2021Ç3–2026Ç2, 20 adet) – ana tablolar',
                       'TBB Ortak Veri Gönderim Seti banka kitabı (31.03.2026 Bin TL, 30.06.2026 Milyon TL) – dipnotlar',
                       'TBB Banka, Şube ve Çalışan Sayıları Haziran 2026 (DT13)', FX_SOURCES['default']],
              gaps=['31.03.2021 ve 30.06.2021 bilançosu KAP bildirimlerinde yok (Önceki Dönem her zaman 31.12)',
                    'Dipnot kalemleri (SYR, LKO, NSFR, ortalama faizler, vade tabloları) yalnızca 31.12.2025 / 31.03.2026 / 30.06.2026 için mevcut (TBB kitapları)',
                    'SYR yalnızca cari dönem için raporlanır: 31.03.2026 ve 30.06.2026',
                    'Hisse başına kâr bankanın bildirdiği gibi; YTD ile çeyrek tutarsız, kullanılmadı']),
    periods=periods,
    fx=dict(rates={p: dict(usd=u, eur=e) for p, (u, e) in FX.items()}, sources=FX_SOURCES,
            method=dict(stock='dönem sonu TCMB döviz alış', flow_q='çeyrek: (önceki dönem sonu + dönem sonu)/2 ortalama kur', flow_ytd='YTD: çeyreklik USD/EUR değerlerinin toplamı')),
    items=items, flows_q=list(flows_q.values()), ratios=ratios, statements=statements, checks=checks, restatements=rest,
    dipnot=dict(scalars=dip_scalars, tables=dip_tables),
)
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
json.dump(out, open(OUT, 'w'), ensure_ascii=False, separators=(',', ':'))
print('written', OUT, os.path.getsize(OUT) // 1024, 'KB', len(items), 'items', len(ratios), 'ratios', len(dip_scalars), 'dipnot scalars')
